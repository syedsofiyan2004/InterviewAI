"""
Builds the MAP TCO Eligibility workbook (Eligibility Checklist + MAP Calculation tabs).

Usage: edit SERVICES and ARR below (or adapt to load from a parsed-data dict), then run.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT_PATH = "/home/claude/map-tco-eligibility/output/Humm_MAP_Eligibility.xlsx"

FONT_NAME = "Arial"
BLUE = Font(name=FONT_NAME, color="0000FF")
BLACK = Font(name=FONT_NAME, color="000000")
BOLD = Font(name=FONT_NAME, bold=True)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="666666")
CURRENCY = '$#,##0;($#,##0);-'
PERCENT = '0.0%'

# --- Parsed service data (from Humm-TCO.pdf) -------------------------------
SERVICES = [
    # name, product_code, category, monthly, eligibility, notes
    ("Amazon Aurora PostgreSQL-Compatible DB", "AmazonRDS", "DB&A", 61741.71,
     "Eligible", "Matched via RDS Aurora PostgreSQL. Counts toward DB&A credit. Confirm if this is an SAP/Oracle workload (would also count toward SAP&Oracle credit)."),
    ("Amazon Elastic Block Store (EBS)", "AmazonEC2", "General", 56389.63,
     "Eligible", "EBS is a sub-feature of EC2 in the included services list, not a separate line."),
    ("Amazon EC2", "AmazonEC2", "General", 25367.50,
     "Partially Eligible", "Excludes Capacity Block for ML (not itemized separately in this TCO export; assume not applicable). Data transfer costs are never eligible and are not broken out here."),
    ("Elastic Load Balancing", "AWSELB", "General", 1592.86,
     "Eligible", ""),
    ("AWS Fargate", "AmazonECS / AmazonEKS", "General", 306.72,
     "Needs Confirmation", "Fargate is eligible under Amazon ECS but NOT under Amazon EKS. TCO export doesn't state which orchestrator is used - confirm with delivery team."),
    ("Amazon Virtual Private Cloud (VPC) - Transit Gateway", "AmazonVPC", "General", 119.60,
     "Eligible", "Config summary shows Transit Gateway attachments; matched to AWS Transit Gateway."),
    ("AWS Security Hub", "AWSSecurityHub", "General", 0.40,
     "Eligible", ""),
]

ARR_TOTAL = round(sum(s[3] for s in SERVICES) * 12, 2)  # should equal PDF's stated $1,746,221.04

def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def build():
    wb = openpyxl.Workbook()

    # ---------------- Tab 1: Eligibility Checklist ----------------
    ws1 = wb.active
    ws1.title = "Eligibility Checklist"
    ws1["A1"] = "AWS MAP Eligibility Checklist"
    ws1["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
    ws1["A2"] = "Estimate only - not a binding MAP/MAP Lite funding determination. Confirm against current AWS Partner Central terms."
    ws1["A2"].font = NOTE_FONT

    headers = ["Service", "Product Code", "Category", "Monthly Cost", "Annual Cost", "Eligibility", "Notes / Exclusions"]
    hdr_row = 4
    for i, h in enumerate(headers, start=1):
        ws1.cell(row=hdr_row, column=i, value=h)
    style_header_row(ws1, hdr_row, len(headers))

    r = hdr_row + 1
    first_data_row = r
    for name, code, cat, monthly, elig, notes in SERVICES:
        ws1.cell(row=r, column=1, value=name).font = BLACK
        ws1.cell(row=r, column=2, value=code).font = BLACK
        ws1.cell(row=r, column=3, value=cat).font = BLACK
        c4 = ws1.cell(row=r, column=4, value=monthly); c4.font = BLUE; c4.number_format = CURRENCY
        c5 = ws1.cell(row=r, column=5, value=f"=D{r}*12"); c5.font = BLACK; c5.number_format = CURRENCY
        ws1.cell(row=r, column=6, value=elig).font = BLACK
        ws1.cell(row=r, column=7, value=notes).font = BLACK
        ws1.cell(row=r, column=7).alignment = Alignment(wrap_text=True)
        r += 1
    last_data_row = r - 1

    total_row = r + 1
    ws1.cell(row=total_row, column=1, value="Total").font = BOLD
    ws1.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{last_data_row})").number_format = CURRENCY
    ws1.cell(row=total_row, column=5, value=f"=SUM(E{first_data_row}:E{last_data_row})").number_format = CURRENCY
    for c in (4, 5):
        ws1.cell(row=total_row, column=c).font = BOLD

    elig_row = total_row + 1
    ws1.cell(row=elig_row, column=1, value="Eligible + Partially Eligible Annual Spend").font = BOLD
    ws1.cell(row=elig_row, column=5,
             value=f'=SUMIFS(E{first_data_row}:E{last_data_row},F{first_data_row}:F{last_data_row},"Eligible")'
                   f'+SUMIFS(E{first_data_row}:E{last_data_row},F{first_data_row}:F{last_data_row},"Partially Eligible")'
             ).number_format = CURRENCY
    ws1.cell(row=elig_row, column=5).font = BOLD

    dba_row = elig_row + 1
    ws1.cell(row=dba_row, column=1, value="DB&A-Eligible Annual Spend").font = BOLD
    ws1.cell(row=dba_row, column=5,
             value=f'=SUMIFS(E{first_data_row}:E{last_data_row},C{first_data_row}:C{last_data_row},"DB&A")'
             ).number_format = CURRENCY

    sap_row = dba_row + 1
    ws1.cell(row=sap_row, column=1, value="SAP&Oracle-Eligible Annual Spend").font = BOLD
    ws1.cell(row=sap_row, column=5,
             value=f'=SUMIFS(E{first_data_row}:E{last_data_row},C{first_data_row}:C{last_data_row},"SAP&Oracle")'
             ).number_format = CURRENCY

    autosize(ws1, [42, 20, 12, 14, 16, 18, 60])
    ws1.freeze_panes = "A5"

    # ---------------- Tab 2: MAP Calculation ----------------
    ws2 = wb.create_sheet("MAP Calculation")
    ws2["A1"] = "AWS MAP Funding Estimate"
    ws2["A1"].font = Font(name=FONT_NAME, bold=True, size=14)
    ws2["A2"] = "Estimate only - not a binding MAP/MAP Lite funding determination. VMware per-VM cash line is internal-only; not for customer visibility."
    ws2["A2"].font = NOTE_FONT

    row = 4
    def label(r, text, bold=True):
        c = ws2.cell(row=r, column=1, value=text)
        c.font = BOLD if bold else BLACK
        return c

    def value_cell(r, col, val, fmt=None, input_cell=False):
        c = ws2.cell(row=r, column=col, value=val)
        c.font = BLUE if input_cell else BLACK
        if input_cell:
            c.fill = INPUT_FILL
        if fmt:
            c.number_format = fmt
        return c

    label(row, "ARR (12-month total, from TCO export)")
    ARR_CELL = f"B{row}"
    value_cell(row, 2, ARR_TOTAL, CURRENCY)  # linked conceptually to Tab1 total; hardcode + note since source is a separate PDF
    ws2.cell(row=row, column=3, value="Source: TCO Calculator 'Total 12 months cost'").font = NOTE_FONT
    row += 1

    label(row, "Eligible ARR (Eligible + Partially Eligible spend)")
    ELIG_ARR_CELL = f"B{row}"
    value_cell(row, 2, "='Eligibility Checklist'!E" + str(elig_row), CURRENCY)
    row += 1

    label(row, "DB&A-eligible ARR")
    DBA_ARR_CELL = f"B{row}"
    value_cell(row, 2, "='Eligibility Checklist'!E" + str(dba_row), CURRENCY)
    row += 1

    label(row, "SAP&Oracle-eligible ARR")
    SAP_ARR_CELL = f"B{row}"
    value_cell(row, 2, "='Eligibility Checklist'!E" + str(sap_row), CURRENCY)
    row += 2

    label(row, "MAP Tier")
    TIER_CELL = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({ARR_CELL}>=500000,{ARR_CELL}<=10000000),"MAP",'
        f'IF(AND({ARR_CELL}>=100000,{ARR_CELL}<500000),"MAP Lite",'
        f'IF(AND({ARR_CELL}>=1000,{ARR_CELL}<100000),"MAP Lite (small)","Not MAP-eligible")))'
    )
    row += 2

    # --- Inputs the user must supply (yellow) ---
    label(row, "INPUTS - fill in per engagement", bold=True)
    row += 1
    label(row, "AWS Greenfield-designated? (Yes/No)")
    GREENFIELD_CELL = f"B{row}"
    value_cell(row, 2, "No", input_cell=True)
    row += 1
    label(row, "% of workloads that are VMware")
    VMWARE_PCT_CELL = f"B{row}"
    value_cell(row, 2, 0, PERCENT, input_cell=True)
    row += 1
    label(row, "% of migration scope that is 'modern services'")
    MODERN_PCT_CELL = f"B{row}"
    value_cell(row, 2, 0, PERCENT, input_cell=True)
    row += 1
    label(row, "Number of VMs being migrated (if VMware modifier claimed)")
    VM_COUNT_CELL = f"B{row}"
    value_cell(row, 2, 0, input_cell=True)
    row += 2

    # --- Assess ---
    label(row, "PHASE 1: ASSESS", bold=True)
    row += 1
    label(row, "Assess cash (5% of ARR, capped at $75,000; MAP / MAP Lite $100K-$500K only)")
    ASSESS_CELL = f"B{row}"
    value_cell(row, 2,
        f'=IF({TIER_CELL}="MAP Lite (small)",0,MIN({ARR_CELL}*0.05,75000))', CURRENCY)
    row += 2

    # --- Mobilize ---
    label(row, "PHASE 2: MOBILIZE", bold=True)
    row += 1
    label(row, "Modernization services (+10% of ARR, up to $100K; ARR must be $500K+)")
    MODERN_CASH = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({TIER_CELL}="MAP",{MODERN_PCT_CELL}>=0.4),MIN({ARR_CELL}*0.1,100000),0)', CURRENCY)
    row += 1
    label(row, "Greenfield (+10% of ARR, up to $100K)")
    GREENFIELD_CASH = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({TIER_CELL}<>"MAP Lite (small)",{GREENFIELD_CELL}="Yes"),MIN({ARR_CELL}*0.1,100000),0)', CURRENCY)
    row += 1
    label(row, "VMware % of ARR (+10% of ARR, up to $200K; needs 75%+ VMware workloads)")
    VMWARE_PCT_CASH = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({TIER_CELL}<>"MAP Lite (small)",{VMWARE_PCT_CELL}>=0.75),MIN({ARR_CELL}*0.1,200000),0)', CURRENCY)
    row += 1
    label(row, "VMware per-VM ($200/VM, max $1M) - INTERNAL ONLY, not for customer visibility")
    VMWARE_VM_CASH = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({TIER_CELL}<>"MAP Lite (small)",{VMWARE_PCT_CELL}>=0.75),MIN({VM_COUNT_CELL}*200,1000000),0)', CURRENCY)
    row += 1
    label(row, "Mobilize cash subtotal (capped at 20% of ARR)")
    MOBILIZE_CELL = f"B{row}"
    value_cell(row, 2,
        f'=MIN(SUM({MODERN_CASH},{GREENFIELD_CASH},{VMWARE_PCT_CASH},{VMWARE_VM_CASH}),{ARR_CELL}*0.2)', CURRENCY)
    row += 2

    # --- Migrate & Modernize ---
    label(row, "PHASE 3: MIGRATE AND MODERNIZE", bold=True)
    row += 1
    label(row, "Base credit (25% of ARR for MAP, 15% for MAP Lite/MAP Lite small)")
    BASE_CREDIT_CELL = f"B{row}"
    value_cell(row, 2,
        f'=IF({TIER_CELL}="MAP",{ARR_CELL}*0.25,IF(OR({TIER_CELL}="MAP Lite",{TIER_CELL}="MAP Lite (small)"),{ARR_CELL}*0.15,0))',
        CURRENCY)
    row += 1
    label(row, "Database & analytics credit (+10% of DB&A-eligible ARR)")
    DBA_CREDIT_CELL = f"B{row}"
    value_cell(row, 2, f'={DBA_ARR_CELL}*0.1', CURRENCY)
    row += 1
    label(row, "SAP & Oracle credit (+50% of SAP&Oracle-eligible ARR)")
    SAP_CREDIT_CELL = f"B{row}"
    value_cell(row, 2, f'={SAP_ARR_CELL}*0.5', CURRENCY)
    row += 1
    label(row, "VMware credit (+10% of ARR, MAP Lite only)")
    VMWARE_CREDIT_CELL = f"B{row}"
    value_cell(row, 2,
        f'=IF(AND({TIER_CELL}="MAP Lite",{VMWARE_PCT_CELL}>=0.75),{ARR_CELL}*0.1,0)', CURRENCY)
    row += 1
    label(row, "Migrate & Modernize credits subtotal")
    MM_CELL = f"B{row}"
    value_cell(row, 2, f'=SUM({BASE_CREDIT_CELL},{DBA_CREDIT_CELL},{SAP_CREDIT_CELL},{VMWARE_CREDIT_CELL})', CURRENCY)
    ws2.cell(row=row, column=2).font = BOLD
    row += 2

    # --- Totals ---
    label(row, "TOTALS", bold=True)
    row += 1
    label(row, "Total estimated partner cash (Assess + Mobilize)")
    value_cell(row, 2, f'=SUM({ASSESS_CELL},{MOBILIZE_CELL})', CURRENCY)
    ws2.cell(row=row, column=2).font = BOLD
    row += 1
    label(row, "Total estimated credits (Migrate & Modernize)")
    value_cell(row, 2, f'={MM_CELL}', CURRENCY)
    ws2.cell(row=row, column=2).font = BOLD

    autosize(ws2, [62, 20, 45])

    import os
    os.makedirs("/home/claude/map-tco-eligibility/output", exist_ok=True)
    wb.save(OUT_PATH)
    print("Saved:", OUT_PATH)

if __name__ == "__main__":
    build()
